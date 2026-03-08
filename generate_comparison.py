#!/usr/bin/env python3
"""Generate comparison.json from XLSX comparison + data.json coordinates."""

import json
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


def parse_mesa_number(mesa_str: str | None) -> int | None:
    """Extract number from 'M123' format."""
    if not mesa_str:
        return None
    match = re.match(r"M(\d+)", str(mesa_str).strip())
    return int(match.group(1)) if match else None


# Geocoded coordinates for mesas not in data.json (new 2026 + missing 2022)
# Obtained via OpenStreetMap Nominatim geocoding for La Matanza addresses
GEOCODED_COORDS: dict[str, tuple[float, float]] = {
    "M1":   (-34.6743875, -58.5596194),
    "M3":   (-34.6375897, -58.5614472),
    "M8":   (-34.6431825, -58.5650632),
    "M9":   (-34.6500849, -58.5525772),
    "M13":  (-34.6512228, -58.5710615),
    "M18":  (-34.6642822, -58.5229964),
    "M25":  (-34.6705902, -58.6043133),
    "M27":  (-34.682646,  -58.543883),
    "M29":  (-34.6570814, -58.5885618),
    "M33":  (-34.6814134, -58.50069),
    "M34":  (-34.6969583, -58.511021),
    "M40":  (-34.6949319, -58.4759238),
    "M42":  (-34.7130795, -58.5187298),
    "M48":  (-34.705552,  -58.5466248),
    "M49":  (-34.7080452, -58.5481697),
    "M50":  (-34.7052429, -58.5389718),
    "M55":  (-34.7323694, -58.5257241),
    "M57":  (-34.7078543, -58.6156304),
    "M60":  (-34.702824,  -58.635295),
    "M64":  (-34.6974883, -58.621676),
    "M70":  (-34.6948026, -58.5749281),
    "M74":  (-34.7162818, -58.583206),
    "M88":  (-34.7507531, -58.5868893),
    "M102": (-34.7379228, -58.6302178),
    "M104": (-34.723929,  -58.5939266),
    "M114": (-34.7552655, -58.6161908),
    "M118": (-34.771266,  -58.6578803),
    "M123": (-34.7721667, -58.6219189),
    "M124": (-34.7721667, -58.6219189),
    "M129": (-34.7990611, -58.6229504),
    "M133": (-34.8527778, -58.6570406),
    "M142": (-34.8591634, -58.6651147),
    "M146": (-34.8686774, -58.6794484),
    "M147": (-34.8914194, -58.6788985),
    "M148": (-34.8534783, -58.6357042),
    "M149": (-34.8932944, -58.6929036),
    "M150": (-34.8197114, -58.6475059),
    "M151": (-34.8273354, -58.6370963),
}


def normalize_estado(raw: str) -> str:
    """Normalize estado to: mantiene | nueva | eliminada."""
    raw_lower = raw.strip().lower()
    if "mantiene" in raw_lower:
        return "mantiene"
    if "nueva" in raw_lower:
        return "nueva"
    if "no es mesa" in raw_lower:
        return "eliminada"
    raise ValueError(f"Unknown estado: {raw}")


def main():
    base = Path(__file__).parent
    xlsx_path = base / "13_mayo_2026" / "comparacion_la_matanza_2022_vs_2026.xlsx"
    data_json_path = base / "pradones" / "webapp" / "public" / "data.json"
    output_path = base / "pradones" / "webapp" / "public" / "comparison.json"

    # Load existing data.json for coordinates
    with open(data_json_path, encoding="utf-8") as f:
        app_data = json.load(f)

    # Build mesa_number → {lat, lng, localidad, zona, zona_nombre} lookup
    mesa_coords: dict[int, dict] = {}
    for mesa in app_data["mesas"]:
        mesa_coords[mesa["mesa"]] = {
            "lat": mesa["lat"],
            "lng": mesa["lng"],
            "localidad": mesa.get("localidad", ""),
            "zona": mesa.get("zona", 0),
            "zona_nombre": mesa.get("zona_nombre", ""),
        }

    # Read COMPLETA sheet from XLSX
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb["COMPLETA"]

    entries = []
    for idx, row in enumerate(
        ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True), start=1
    ):
        establecimiento = str(row[0] or "").strip()
        direccion = str(row[1] or "").strip()
        mesa_2022_str = str(row[2] or "").strip() if row[2] else None
        mesa_2026_str = str(row[3] or "").strip() if row[3] else None
        estado_raw = str(row[4] or "").strip()

        if not establecimiento and not estado_raw:
            continue

        estado = normalize_estado(estado_raw)
        mesa_2022 = parse_mesa_number(mesa_2022_str)
        mesa_2026 = parse_mesa_number(mesa_2026_str)

        # Get coordinates from data.json via 2022 mesa number
        coords = None
        if mesa_2022 and mesa_2022 in mesa_coords:
            coords = mesa_coords[mesa_2022]

        lat = coords["lat"] if coords else None
        lng = coords["lng"] if coords else None

        # Fallback: use geocoded coordinates for mesas not in data.json
        if lat is None and mesa_2026_str and mesa_2026_str in GEOCODED_COORDS:
            lat, lng = GEOCODED_COORDS[mesa_2026_str]

        entry = {
            "id": idx,
            "establecimiento": establecimiento,
            "direccion": direccion,
            "mesa_2022": mesa_2022_str,
            "mesa_2026": mesa_2026_str,
            "estado": estado,
            "lat": lat,
            "lng": lng,
            "localidad": coords["localidad"] if coords else None,
            "zona_2022": coords["zona"] if coords else None,
            "zona_nombre_2022": coords["zona_nombre"] if coords else None,
        }
        entries.append(entry)

    # Also read extra info from SE MANTIENEN (match type, changed number)
    ws_mantiene = wb["SE MANTIENEN"]
    mantiene_extra: dict[str, dict] = {}
    for row in ws_mantiene.iter_rows(min_row=2, max_row=ws_mantiene.max_row, values_only=True):
        estab = str(row[0] or "").strip()
        cambio = str(row[4] or "").strip() if row[4] else "NO"
        match_por = str(row[5] or "").strip() if row[5] else ""
        mantiene_extra[estab] = {"cambio_mesa": cambio != "NO", "match_por": match_por}

    for entry in entries:
        if entry["estado"] == "mantiene" and entry["establecimiento"] in mantiene_extra:
            extra = mantiene_extra[entry["establecimiento"]]
            entry["cambio_mesa"] = extra["cambio_mesa"]
            entry["match_por"] = extra["match_por"]

    # Read extra info from YA NO SON MESA (zona 2022)
    ws_removed = wb["YA NO SON MESA"]
    for row in ws_removed.iter_rows(min_row=2, max_row=ws_removed.max_row, values_only=True):
        estab = str(row[1] or "").strip()
        localidad = str(row[3] or "").strip() if row[3] else None
        zona_str = str(row[4] or "").strip() if row[4] else None
        for entry in entries:
            if entry["estado"] == "eliminada" and entry["establecimiento"] == estab:
                if localidad and not entry.get("localidad"):
                    entry["localidad"] = localidad
                if zona_str and not entry.get("zona_nombre_2022"):
                    entry["zona_nombre_2022"] = zona_str
                break

    # Summary stats
    total_2022 = 158
    total_2026 = 151
    maintained = sum(1 for e in entries if e["estado"] == "mantiene")
    new_2026 = sum(1 for e in entries if e["estado"] == "nueva")
    removed = sum(1 for e in entries if e["estado"] == "eliminada")
    with_coords = sum(1 for e in entries if e["lat"] is not None)

    summary = {
        "total_2022": total_2022,
        "total_2026": total_2026,
        "maintained": maintained,
        "new_2026": new_2026,
        "removed": removed,
        "with_coords": with_coords,
    }

    result = {
        "generated": "2026-03-02",
        "summary": summary,
        "entries": entries,
    }

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"Generated {output_path}")
    print(f"  Total entries: {len(entries)}")
    print(f"  Maintained: {maintained}, New: {new_2026}, Removed: {removed}")
    print(f"  With coordinates: {with_coords}")


def generate_xlsx():
    """Generate a downloadable XLSX with all 2026 mesas including geolocation."""
    base = Path(__file__).parent
    comparison_path = base / "pradones" / "webapp" / "public" / "comparison.json"
    output_path = base / "pradones" / "webapp" / "public" / "mesas_2026_la_matanza.xlsx"

    with open(comparison_path, encoding="utf-8") as f:
        data = json.load(f)

    # Filter only entries that have a mesa_2026 (mantiene + nueva)
    entries_2026 = [e for e in data["entries"] if e.get("mesa_2026")]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mesas 2026 La Matanza"

    headers = [
        "Mesa 2026",
        "Establecimiento",
        "Dirección",
        "Estado",
        "Mesa 2022",
        "Latitud",
        "Longitud",
        "Google Maps",
    ]
    header_font = Font(bold=True)
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, entry in enumerate(entries_2026, start=2):
        lat = entry.get("lat")
        lng = entry.get("lng")

        ws.cell(row=row_idx, column=1, value=entry["mesa_2026"])
        ws.cell(row=row_idx, column=2, value=entry["establecimiento"])
        ws.cell(row=row_idx, column=3, value=entry["direccion"])
        ws.cell(row=row_idx, column=4, value=entry["estado"])
        ws.cell(row=row_idx, column=5, value=entry.get("mesa_2022", ""))
        ws.cell(row=row_idx, column=6, value=lat)
        ws.cell(row=row_idx, column=7, value=lng)

        if lat is not None and lng is not None:
            maps_url = f"https://www.google.com/maps?q={lat},{lng}"
            cell = ws.cell(row=row_idx, column=8, value=maps_url)
            cell.hyperlink = maps_url
            cell.font = Font(color="0563C1", underline="single")
        else:
            ws.cell(row=row_idx, column=8, value="Sin coordenadas")

    # Auto-width columns
    for col_idx in range(1, len(headers) + 1):
        max_len = len(str(headers[col_idx - 1]))
        for row in ws.iter_rows(
            min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx
        ):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)

    wb.save(output_path)
    print(f"Generated XLSX: {output_path}")
    print(f"  Mesas 2026: {len(entries_2026)}")


if __name__ == "__main__":
    main()
    generate_xlsx()
